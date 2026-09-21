from fastapi import APIRouter, UploadFile, File, HTTPException, Form
from datetime import datetime
import uuid
import io

from services.supabase_client import supabase_admin
from models.rag_engine import RAGEngine
from models.embeddings import EmbeddingService

router = APIRouter(prefix="/api/documents", tags=["documents"])

rag_engine = RAGEngine()
embedder = EmbeddingService()


def _sanitize_text(text: str) -> str:
    """Remove null bytes and control chars that Postgres rejects."""
    if not text:
        return ""
    text = text.replace("\x00", "").replace("\u0000", "")
    text = "".join(
        ch for ch in text
        if ch in ("\n", "\r", "\t") or (ord(ch) >= 32 and ord(ch) != 127)
    )
    return text


def _extract_text_from_pdf(content: bytes) -> str:
    """Extract text from PDF using pdfplumber (fallback: PyPDF2)."""
    print(f"[_extract_text_from_pdf] Attempting pdfplumber on {len(content)} bytes")
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = []
            for page in pdf.pages:
                txt = page.extract_text() or ""
                pages.append(txt)
            return "\n\n".join(pages).strip()
    except Exception as e1:
        print(f"[PDF] pdfplumber failed: {e1}. Trying PyPDF2...")
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            return "\n\n".join(pages).strip()
        except Exception as e2:
            raise Exception(f"PDF extraction failed: {e1} / {e2}")


def _extract_text_from_docx(content: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        # Also grab table text
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        table_text.append(cell.text)
        all_text = "\n".join(paragraphs + table_text)
        return all_text.strip()
    except Exception as e:
        raise Exception(f"DOCX extraction failed: {e}")


def _extract_text(content: bytes, filename: str, content_type: str = None) -> str:
    """Route extraction based on file type."""
    fname = (filename or "").lower()
    ctype = (content_type or "").lower()
    
    print(f"[_extract_text] fname={fname!r} ctype={ctype!r} size={len(content)} first20={content[:20]!r}")

    is_pdf = fname.endswith(".pdf") or "pdf" in ctype
    is_docx = fname.endswith(".docx") or "wordprocessingml" in ctype or "msword" in ctype
    is_doc = fname.endswith(".doc") and not is_docx

    if is_pdf:
        text = _extract_text_from_pdf(content)
        if not text.strip():
            raise Exception("PDF has no extractable text (might be scanned/image-only).")
        return text

    if is_docx:
        text = _extract_text_from_docx(content)
        if not text.strip():
            raise Exception("DOCX has no readable text.")
        return text

    if is_doc:
        raise Exception("Legacy .doc format not supported. Please save as .docx or .pdf.")

    # Default: treat as text file
    for encoding in ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "latin-1"):
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("utf-8", errors="ignore")


@router.get("")
async def list_documents(user_id: str = "default-user"):
    res = (supabase_admin.table("documents")
           .select("*")
           .eq("user_id", user_id)
           .order("created_at", desc=True)
           .execute())
    return {"documents": res.data, "total": len(res.data), "status": "success"}


@router.get("/{document_id}")
async def get_document(document_id: str):
    res = supabase_admin.table("documents").select("*").eq("id", document_id).execute()
    if not res.data:
        raise HTTPException(404, "Document not found")
    return {"status": "success", "document": res.data[0]}


@router.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
    user_id: str = Form("default-user"),
):
    try:
        content = await file.read()
        safe_filename = _sanitize_text(file.filename or "untitled")[:500]

        # Safety cap: prevent OOM on free tier
        MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
        if len(content) > MAX_FILE_SIZE:
            raise HTTPException(
                413,
                f"File too large ({len(content) / 1024 / 1024:.1f} MB). Max 10 MB."
            )

        # ===== Extract text based on file type =====
        try:
            text_content = _extract_text(content, file.filename, file.content_type)
        except Exception as extract_err:
            raise HTTPException(400, f"Could not extract text: {extract_err}")

        # ===== Sanitize text =====
        text_content = _sanitize_text(text_content)

        if not text_content.strip():
            raise HTTPException(400, "No readable text found in the document.")

        doc_id = f"doc_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}"

        supabase_admin.table("documents").insert({
            "id": doc_id,
            "user_id": user_id,
            "name": safe_filename,
            "file_type": file.content_type or "text/plain",
            "file_size": len(content),
            "content": text_content[:100000],
            "status": "processing",
            "document_type": "legal",
        }).execute()

        # Chunk & embed
        chunks = rag_engine.chunk_text(text_content)
        rows = []
        for i, chunk in enumerate(chunks):
            safe_chunk = _sanitize_text(chunk)
            if not safe_chunk.strip():
                continue
            rows.append({
                "document_id": doc_id,
                "chunk_index": i,
                "text": safe_chunk,
                "token_count": len(safe_chunk.split()),
                "embedding": embedder.encode(safe_chunk),
            })

        if rows:
            supabase_admin.table("document_chunks").insert(rows).execute()

        supabase_admin.table("documents").update({
            "status": "analyzed",
            "total_chunks": len(rows),
            "last_analyzed_at": datetime.utcnow().isoformat(),
        }).eq("id", doc_id).execute()

        return {
            "status": "success",
            "document_id": doc_id,
            "filename": safe_filename,
            "total_chunks": len(rows),
            "text_length": len(text_content),
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, str(e))


@router.delete("/{document_id}")
async def delete_document(document_id: str):
    try:
        supabase_admin.table("document_chunks").delete().eq("document_id", document_id).execute()
        supabase_admin.table("documents").delete().eq("id", document_id).execute()
        return {"status": "success", "message": "Document deleted", "document_id": document_id}
    except Exception as e:
        raise HTTPException(500, str(e))